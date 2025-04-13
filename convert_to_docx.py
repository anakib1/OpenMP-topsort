import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import markdown
import re

def create_docx_from_markdown(md_file, docx_file):
    # Create a new document
    doc = docx.Document()

    # Read markdown content
    with open(md_file, 'r', encoding='utf-8') as file:
        md_content = file.read()

    # Split content into sections
    sections = md_content.split('\n## ')
    
    # Process title
    title = sections[0].split('\n')[0].replace('# ', '')
    heading = doc.add_heading(title, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Process each section
    for section in sections[1:]:
        lines = section.split('\n')
        section_title = lines[0]
        doc.add_heading(section_title, 1)
        
        content = '\n'.join(lines[1:])
        
        # Process tables
        if '|' in content:
            table_lines = [line for line in content.split('\n') if '|' in line]
            if len(table_lines) > 2:  # Has header and content
                # Get number of columns
                cols = len(table_lines[0].split('|')) - 2
                # Create table
                table = doc.add_table(rows=1, cols=cols)
                table.style = 'Table Grid'
                
                # Add header
                header_cells = table.rows[0].cells
                headers = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]
                for i, text in enumerate(headers):
                    header_cells[i].text = text
                
                # Add data rows
                for row in table_lines[2:]:  # Skip header and separator
                    cells = [cell.strip() for cell in row.split('|')[1:-1]]
                    row_cells = table.add_row().cells
                    for i, text in enumerate(cells):
                        row_cells[i].text = text
                
                doc.add_paragraph()  # Add space after table
        
        # Process code blocks
        code_blocks = re.findall(r'```.*?\n(.*?)```', content, re.DOTALL)
        for block in code_blocks:
            p = doc.add_paragraph()
            run = p.add_run(block)
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            content = content.replace(f'```cpp\n{block}```', '')
        
        # Process remaining content
        paragraphs = [p for p in content.split('\n') if p and not p.startswith('```')]
        for para in paragraphs:
            if para.strip():
                if para.startswith('- '):
                    doc.add_paragraph(para[2:], style='List Bullet')
                elif para.startswith('1. '):
                    doc.add_paragraph(para[3:], style='List Number')
                else:
                    doc.add_paragraph(para)

    # Save the document
    doc.save(docx_file)

if __name__ == '__main__':
    create_docx_from_markdown('report.md', 'report.docx') 